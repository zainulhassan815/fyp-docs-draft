#!/usr/bin/env python3
"""Make all tables auto-fit to window width in a Word document."""

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt
import sys

# Width of the text column, in twips (Letter page less the SCET template margins).
CONTENT_TWIPS = 8640

# Most tables are fine with the equal columns that autofit gives them, because
# the lua filter strips pandoc's width hints.  A few carry cells whose content
# is far more uneven than that, and equal columns break long tokens mid-word.
# Give those explicit proportions, keyed by their header row.
COLUMN_WIDTHS = {
    ('Test Case', 'Input', 'Expected Result', 'Actual Result', 'Status'):
        (26, 20, 23, 21, 10),
    ('Use Case ID', 'Use Case Name', 'Functional Requirement ID',
     'Test Case ID(s)', 'Test Scenario', 'Status'):
        (10, 16, 16, 24, 25, 9),
    ('#', 'Stage', 'Input', 'Processing', 'Output'):
        (5, 14, 17, 38, 26),
    ('Ref', 'System (Year)', 'Research Problem', 'Methodology & Tools',
     'Key Features', 'Limitations', 'Relevance to Hireflow'):
        (7, 14, 15, 19, 14, 15, 16),
}


def header_key(table):
    return tuple(cell.text.strip() for cell in table.rows[0].cells)


def set_column_widths(table, weights):
    """Pin proportional column widths; autofit cannot infer them without hints."""
    total = sum(weights)
    widths = [round(CONTENT_TWIPS * w / total) for w in weights]

    tblPr = table._tbl.tblPr
    for child in tblPr.findall(qn('w:tblLayout')):
        tblPr.remove(child)
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)

    grid = table._tbl.find(qn('w:tblGrid'))
    if grid is not None:
        for col, width in zip(grid.findall(qn('w:gridCol')), widths):
            col.set(qn('w:w'), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Pt(width / 20)


def set_table_font(table, pt):
    """Set the font of every run in a table so wide tables fit the page width."""
    size = Pt(pt)
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = size


def font_pt_for_columns(n_cols):
    """Graduated font size: only shrink genuinely wide tables, leave narrow ones alone."""
    if n_cols >= 8:
        return 8
    if n_cols >= 6:
        return 9
    if n_cols == 5:
        return 10
    return None  # <= 4 columns: keep the document default size


def set_table_autofit_window(table):
    """Set a table to auto-fit to window width."""
    tbl = table._tbl

    # Get or create tblPr element
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)

    # Remove existing width settings
    for child in tblPr.findall(qn('w:tblW')):
        tblPr.remove(child)

    # Set table width to 100% (5000 = 100% in fiftieths of a percent)
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), '5000')
    tblW.set(qn('w:type'), 'pct')
    tblPr.append(tblW)

    # Remove fixed layout if present, set to autofit
    for child in tblPr.findall(qn('w:tblLayout')):
        tblPr.remove(child)

    tblLayout = OxmlElement('w:tblLayout')
    tblLayout.set(qn('w:type'), 'autofit')
    tblPr.append(tblLayout)


def autofit_tables(docx_path, output_path=None):
    """Process all tables in a document to auto-fit to window."""
    if output_path is None:
        output_path = docx_path

    doc = Document(docx_path)
    table_count = len(doc.tables)

    for table in doc.tables:
        set_table_autofit_window(table)
        pt = font_pt_for_columns(len(table.columns))
        if pt is not None:
            set_table_font(table, pt)
        weights = COLUMN_WIDTHS.get(header_key(table))
        if weights and len(weights) == len(table.columns):
            set_column_widths(table, weights)

    doc.save(output_path)
    print(f"Processed {table_count} table(s) in {docx_path}")


def main():
    if len(sys.argv) < 2:
        # Default: process the SRS document
        docx_path = 'build/SRS_Document.docx'
    else:
        docx_path = sys.argv[1]

    output_path = sys.argv[2] if len(sys.argv) > 2 else None
    autofit_tables(docx_path, output_path)


if __name__ == '__main__':
    main()
