#!/usr/bin/env python3
"""Generate the hardbound book-cover artwork for the print shop.

Follows the layout the department's bound reports use: title, rule, crest,
supervisors, the submitting students, rule, then the institution block.

Everything is set in black on white. The shop foils it gold; supplying the
artwork in colour would only confuse the block-making. Margins are deliberately
generous so nothing sits near the board edge, where a case-bound cover curves.

Output: build/Book_Cover.docx
"""

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

OUT = 'build/Book_Cover.docx'
CREST = 'src/images/scet-crest.jpeg'
FONT = 'Times New Roman'

TITLE = ['AI POWERED HR SCREENING AND DOCUMENT',
         'RETRIEVAL SYSTEM USING RAG']
SUPERVISORS = ['SUPERVISOR: MS. HIRRA MUSTAFA',
               'CO-SUPERVISOR: DR. MAZHAR IQBAL']
STUDENTS = [('AMNA IKRAM', '2022-UET-SHCET-LHR-CS-02'),
            ('ZAIN UL HASSAN', '2022-UET-SHCET-LHR-CS-12'),
            ('EZZA ANSAR', '2022-UET-SHCET-LHR-CS-13')]
FOOT = ['DEPARTMENT OF COMPUTER SCIENCE',
        'SHARIF COLLEGE OF ENGINEERING & TECHNOLOGY',
        'LAHORE - PAKISTAN',
        'SEPTEMBER, 2026']


def line(doc, text, size, *, bold=True, before=0, after=0, spacing=1.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), spacing
    run = p.add_run(text)
    run.bold = bold
    run.font.name = FONT
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn('w:cs'), FONT)
    return p


def rule(doc, *, before=0, after=0, weight=18):
    """A horizontal rule, drawn as a paragraph border."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    borders = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    for key, value in (('val', 'single'), ('sz', str(weight)),
                       ('space', '1'), ('color', '000000')):
        bottom.set(qn('w:' + key), value)
    borders.append(bottom)
    p._p.get_or_add_pPr().append(borders)
    return p


def main():
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)          # A4
    s.top_margin = s.bottom_margin = Inches(1.3)
    s.left_margin = s.right_margin = Inches(1.4)
    s.header_distance = s.footer_distance = Inches(0.8)

    for text in TITLE:
        line(doc, text, 17, spacing=1.25)

    rule(doc, before=10, after=0, weight=18)

    crest = doc.add_paragraph()
    crest.alignment = WD_ALIGN_PARAGRAPH.CENTER
    crest.paragraph_format.space_before = Pt(64)
    crest.paragraph_format.space_after = Pt(64)
    crest.add_run().add_picture(CREST, height=Inches(1.6))

    for n, text in enumerate(SUPERVISORS):
        line(doc, text, 14, before=0 if n else 6, after=2)

    line(doc, 'SUBMITTED BY', 14, before=46, after=8)

    table = doc.add_table(rows=len(STUDENTS), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row, (name, reg) in zip(table.rows, STUDENTS):
        for cell, text, align in ((row.cells[0], name, WD_ALIGN_PARAGRAPH.LEFT),
                                  (row.cells[1], reg, WD_ALIGN_PARAGRAPH.RIGHT)):
            cell.width = Inches(2.6)
            p = cell.paragraphs[0]
            p.alignment = align
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(text)
            run.bold = True
            run.font.name = FONT
            run.font.size = Pt(13)

    rule(doc, before=58, after=12, weight=12)

    for n, text in enumerate(FOOT):
        line(doc, text, 14 if n < 2 else 13, before=0 if n else 2, after=3)

    doc.save(OUT)
    print(f"Wrote {OUT}: A4, margins 1.3in top/bottom, 1.4in sides")


if __name__ == '__main__':
    main()
