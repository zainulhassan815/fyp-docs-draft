#!/usr/bin/env python3
"""Set justified alignment on body-text paragraphs only.

Headings, captions, code blocks, list items, and table-cell paragraphs are left
untouched (table cells are never in Document.paragraphs; the other cases are excluded
by style name)."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
import sys

# Only genuine running-prose body styles are justified.
BODY_STYLES = {"Normal", "Body Text", "First Paragraph"}


def justify(docx_path):
    doc = Document(docx_path)
    count = 0
    for p in doc.paragraphs:
        name = p.style.name if p.style is not None else "Normal"
        # Only justify running prose that has no explicit alignment. This leaves the
        # cover's centred title block (CENTER) and pre-justified declaration untouched.
        if name in BODY_STYLES and p.alignment is None:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            count += 1
    doc.save(docx_path)
    print(f"Justified {count} body paragraph(s) in {docx_path}")


def main():
    justify(sys.argv[1] if len(sys.argv) > 1 else "build/SRS_Document.docx")


if __name__ == "__main__":
    main()
