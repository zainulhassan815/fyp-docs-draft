#!/usr/bin/env python3
"""Merge cover page with document and apply reference styles."""

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docxcompose.composer import Composer
import zipfile
import os
import shutil

def main():
    cover = Document('templates/cover.docx')
    content = Document('build/SRS_Document.docx')

    # Add page break at end of cover
    p = cover.add_paragraph()
    run = p.add_run()
    br = OxmlElement('w:br')
    br.set(qn('w:type'), 'page')
    run._r.append(br)

    # Compose with cover as base (stable)
    composer = Composer(cover)
    composer.append(content)
    composer.save('build/SRS_Document.docx')

    # Replace styles.xml with reference-doc styles
    temp_dir = 'build/temp_docx'
    with zipfile.ZipFile('templates/custom-reference.docx', 'r') as ref:
        styles_xml = ref.read('word/styles.xml')

    with zipfile.ZipFile('build/SRS_Document.docx', 'r') as z:
        z.extractall(temp_dir)

    with open(os.path.join(temp_dir, 'word', 'styles.xml'), 'wb') as f:
        f.write(styles_xml)

    os.remove('build/SRS_Document.docx')
    with zipfile.ZipFile('build/SRS_Document.docx', 'w', zipfile.ZIP_DEFLATED) as z:
        for root, dirs, files in os.walk(temp_dir):
            for file in files:
                path = os.path.join(root, file)
                z.write(path, os.path.relpath(path, temp_dir))

    shutil.rmtree(temp_dir)
    set_a4('build/SRS_Document.docx')


def set_a4(path):
    """Switch the page size to A4, leaving the margins as the template has them.

    The composed document inherits its geometry from cover.docx, which came from
    the SCET template as US Letter. The bound report is printed on A4, so laying
    it out on Letter would reflow on the way to the printer. Margins are not
    touched. A4 is 11906 x 16838 twips.
    """
    import re

    with zipfile.ZipFile(path) as z:
        doc = z.read('word/document.xml').decode('utf-8')
    doc = re.sub(r'<w:pgSz[^>]*/>', '<w:pgSz w:w="11906" w:h="16838"/>', doc)

    tmp = path + '.tmp'
    with zipfile.ZipFile(path) as zin, \
            zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == 'word/document.xml':
                data = doc.encode('utf-8')
            zout.writestr(item, data)
    os.replace(tmp, path)

if __name__ == '__main__':
    main()
