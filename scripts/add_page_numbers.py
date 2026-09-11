#!/usr/bin/env python3
"""Build the footer: department on the left, page number on the right.

The department wants "Department of Computer Science" at the start of the
footer and the page number at the end. One paragraph carries both, separated
by a right-aligned tab stop at the right margin, so the two sit at opposite
edges on every page whatever the page number's width.

A thin rule sits above the footer to match the one under the header, so the
top and bottom edges of a page are balanced.

Numbering follows the usual convention for a bound report: the front matter
runs in lower-case roman numerals and the body restarts at 1 on Chapter 1.
The title page is counted as i but shows no header or footer of its own.
"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt

DOC = 'build/SRS_Document.docx'
DEPARTMENT = 'Department of Computer Science, SCET'
# The SCET template ships w:footer="0", which puts the footer hard against the
# paper edge, inside the strip most printers cannot reach. Word's own default.
FOOTER_DISTANCE = Inches(0.5)
# Matches the rule under the header.
RULE_COLOUR = 'A6A6A6'


def page_field(paragraph):
    """Append a { PAGE } field so Word renders the live page number."""
    run = paragraph.add_run()
    for kind, text in (('begin', None), (None, 'PAGE'), ('end', None)):
        if kind:
            fld = OxmlElement('w:fldChar')
            fld.set(qn('w:fldCharType'), kind)
        else:
            fld = OxmlElement('w:instrText')
            fld.set(qn('xml:space'), 'preserve')
            fld.text = f' {text} '
        run._r.append(fld)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11)


def rule_above(paragraph):
    """Thin rule over the footer, mirroring the one under the header."""
    pPr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement('w:pBdr')
    top = OxmlElement('w:top')
    for key, value in (('val', 'single'), ('sz', '4'), ('space', '6'),
                       ('color', RULE_COLOUR)):
        top.set(qn('w:' + key), value)
    borders.append(top)
    pPr.append(borders)


def page_numbering(section, fmt, start=None):
    """Set the numbering format, and optionally restart the count."""
    sectPr = section._sectPr
    el = sectPr.find(qn('w:pgNumType'))
    if el is None:
        el = OxmlElement('w:pgNumType')
        sectPr.append(el)
    el.set(qn('w:fmt'), fmt)
    if start is None:
        if el.get(qn('w:start')) is not None:
            del el.attrib[qn('w:start')]
    else:
        el.set(qn('w:start'), str(start))


def blank_first_page(section):
    """Give the title page its own empty header and footer."""
    section.different_first_page_header_footer = True
    for part in (section.first_page_header, section.first_page_footer):
        part.is_linked_to_previous = False
        for para in part.paragraphs:
            para.clear()


def main():
    doc = Document(DOC)
    # Front matter in roman, body restarting at 1. Section 0 is the front
    # matter and section 1 the body; add_header.py made that split.
    page_numbering(doc.sections[0], 'lowerRoman')
    if len(doc.sections) > 1:
        page_numbering(doc.sections[1], 'decimal', start=1)

    for section in doc.sections:
        section.different_first_page_header_footer = False
        section.footer_distance = FOOTER_DISTANCE
        footer = section.footer
        footer.is_linked_to_previous = False
        para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        for run in list(para.runs):
            run._r.getparent().remove(run._r)
        para.text = ''
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT

        # One tab stop at the right margin pushes the page number to the edge.
        width = section.page_width - section.left_margin - section.right_margin
        para.paragraph_format.tab_stops.clear_all()
        para.paragraph_format.tab_stops.add_tab_stop(Emu(width), WD_TAB_ALIGNMENT.RIGHT)

        left = para.add_run(DEPARTMENT)
        left.font.name = 'Times New Roman'
        left.font.size = Pt(11)
        para.add_run().add_tab()
        page_field(para)
        rule_above(para)

    # The title page carries neither, but still counts as i.
    blank_first_page(doc.sections[0])

    doc.save(DOC)
    print(f"Footer: '{DEPARTMENT}' left, page number right; "
          f"front matter roman, body restarts at 1")


if __name__ == '__main__':
    main()
